from __future__ import annotations

import json
import os
import time
from dataclasses import dataclass
from pathlib import Path
from typing import List

import gspread
import requests
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt
from dotenv import load_dotenv
from google import genai
from google.oauth2.service_account import Credentials

load_dotenv()

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive.readonly",
]


@dataclass
class Settings:
    gemini_api_key: str
    google_service_account_json: str
    sheet_url: str
    sheet_name: str
    output_dir: Path
    tw_url_col: int = 0
    hk_url_col: int = 1
    start_row: int = 1
    section_char_limit: int = 3500
    request_timeout: int = 20
    request_sleep_seconds: float = 1.0


@dataclass
class Article:
    slug: str
    title: str
    meta_description: str
    content_html: str
    url: str


def load_settings() -> Settings:
    return Settings(
        gemini_api_key=require_env("GEMINI_API_KEY"),
        google_service_account_json=require_env("GOOGLE_SERVICE_ACCOUNT_JSON"),
        sheet_url=require_env("SHEET_URL"),
        sheet_name=os.getenv("SHEET_NAME", "工作表1"),
        output_dir=Path(os.getenv("OUTPUT_DIR", "./output/translated_articles")),
        tw_url_col=int(os.getenv("TW_URL_COL", "0")),
        hk_url_col=int(os.getenv("HK_URL_COL", "1")),
        start_row=int(os.getenv("START_ROW", "1")),
        section_char_limit=int(os.getenv("SECTION_CHAR_LIMIT", "3500")),
        request_timeout=int(os.getenv("REQUEST_TIMEOUT", "20")),
        request_sleep_seconds=float(os.getenv("REQUEST_SLEEP_SECONDS", "1")),
    )


def require_env(key: str) -> str:
    value = os.getenv(key, "").strip()
    if not value:
        raise ValueError(f"Missing required environment variable: {key}")
    return value


def build_gspread_client(settings: Settings) -> gspread.Client:
    credentials = Credentials.from_service_account_file(
        settings.google_service_account_json,
        scopes=SCOPES,
    )
    return gspread.authorize(credentials)


def get_pending_urls(sheet: gspread.Worksheet, settings: Settings) -> List[str]:
    rows = sheet.get_all_values()
    pending_urls = []
    for index, row in enumerate(rows):
        if index < settings.start_row:
            continue
        tw_url = row[settings.tw_url_col].strip() if len(row) > settings.tw_url_col else ""
        hk_url = row[settings.hk_url_col].strip() if len(row) > settings.hk_url_col else ""
        if tw_url and not hk_url:
            pending_urls.append(tw_url)
    return pending_urls


def fetch_article(url: str, timeout: int) -> Article | None:
    headers = {"User-Agent": "Mozilla/5.0"}
    response = requests.get(url, headers=headers, timeout=timeout)
    response.raise_for_status()
    response.encoding = response.apparent_encoding or "utf-8"
    soup = BeautifulSoup(response.text, "html.parser")

    title_tag = soup.select_one("h1.entry-title") or soup.find("h1")
    meta_tag = soup.find("meta", attrs={"name": "description"}) or soup.find(
        "meta",
        attrs={"property": "og:description"},
    )
    content_div = soup.select_one("div.entry-content")
    if not content_div:
        return None

    cleaned_html = clean_content_html(content_div)
    slug = url.rstrip("/").split("/")[-1]
    return Article(
        slug=slug,
        title=title_tag.get_text(strip=True) if title_tag else "",
        meta_description=meta_tag.get("content", "") if meta_tag else "",
        content_html=cleaned_html,
        url=url,
    )


def clean_content_html(content_div: Tag) -> str:
    cloned = BeautifulSoup(str(content_div), "html.parser")
    root = cloned.select_one("div.entry-content") or cloned

    for style_tag in root.find_all("style"):
        style_tag.decompose()

    for removable in root.select(".share-buttons, .slider, .product-section"):
        removable.decompose()

    for gap in root.select("div.gap-element"):
        gap.replace_with(cloned.new_tag("gap"))

    unwrap_classes = {"row", "col", "col-inner", "container"}
    for div in list(root.find_all("div")):
        class_names = set(div.get("class", []))
        if class_names & unwrap_classes and not class_names.intersection({"text", "section-title-container"}):
            div.unwrap()

    return str(root)


def split_html_into_sections(html: str, char_limit: int) -> List[str]:
    soup = BeautifulSoup(html, "html.parser")
    root = soup.select_one("div.entry-content") or soup
    sections: List[str] = []
    current: List[str] = []

    def flush() -> None:
        nonlocal current
        if current:
            sections.append("".join(current).strip())
            current = []

    for child in root.children:
        if isinstance(child, NavigableString) and not child.strip():
            continue
        child_html = str(child)
        child_name = getattr(child, "name", "")
        if child_name in {"h2", "h3", "h4"} and current:
            flush()
        if sum(len(part) for part in current) + len(child_html) > char_limit and current:
            flush()
        current.append(child_html)
    flush()

    if not sections:
        return [html]
    return sections


def translate_text(client: genai.Client, text: str, content_type: str) -> str:
    if not text.strip():
        return ""

    if content_type == "body_html":
        prompt = f"""
你是一位香港 SEO 在地化編輯。
請將以下台灣繁體中文 HTML 內容，改寫成自然香港粵語。
規則：
- 保留原有 HTML 標籤與結構
- 不可修改屬性
- 只改標籤之間文字
- 不可新增說明
- 用自然香港用語，但不要過度擴寫
HTML：
{text}
""".strip()
    else:
        prompt = f"""
請把以下內容改成自然香港粵語，不要加說明。
內容：
{text}
""".strip()

    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=prompt,
    )
    result = (response.text or "").strip()
    if not result:
        raise ValueError(f"Empty model output for {content_type}")
    return result


def convert_html_to_shortcodes(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    output: List[str] = []

    for node in soup.contents:
        converted = convert_node(node)
        if converted:
            output.append(converted)
    return "\n\n".join(part.strip() for part in output if part.strip())


def convert_node(node) -> str:
    if isinstance(node, NavigableString):
        return str(node).strip()
    if not isinstance(node, Tag):
        return ""

    class_names = set(node.get("class", []))
    title_span = node.select_one(".section-title-main")
    heading = node.find(["h2", "h3", "h4"])

    if node.name == "gap":
        return "[gap]"

    if "section-title-container" in class_names and title_span and heading:
        text = title_span.get_text(strip=True).replace('"', "'")
        tag_name = heading.name
        if tag_name == "h2":
            return f'[title style="bold-center" text="{text}" tag_name="h2" color="rgb(10, 184, 188)" size="120"]'
        if tag_name == "h4":
            return f'[title text="{text}" tag_name="h4" color="rgb(10, 184, 188)" size="120"]'
        return f'[title text="{text}" color="rgb(10, 184, 188)" size="120"]'

    if node.name == "div" and "text" in class_names:
        inner = "".join(convert_node(child) if isinstance(child, Tag) else str(child) for child in node.contents).strip()
        return f'[ux_text line_height="2"]{inner}[/ux_text]'

    if node.name == "div":
        return "\n".join(filter(None, (convert_node(child) for child in node.contents)))

    return str(node)


def add_field(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    label_run = paragraph.add_run(label)
    label_run.bold = True
    label_run.font.size = Pt(12)
    value_run = paragraph.add_run(value)
    value_run.font.size = Pt(12)
    paragraph.paragraph_format.space_after = Pt(8)


def create_docx(article: Article, translated_title: str, translated_meta: str, shortcode_content: str, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(12)

    add_field(doc, "標題：", translated_title)
    add_field(doc, "描述：", translated_meta)
    add_field(doc, "網址：", article.slug)
    doc.add_paragraph("─" * 50)
    body_label = doc.add_paragraph()
    run = body_label.add_run("內文：")
    run.bold = True
    run.font.size = Pt(12)
    body_paragraph = doc.add_paragraph(shortcode_content)
    if body_paragraph.runs:
        body_paragraph.runs[0].font.size = Pt(11)

    file_path = output_dir / f"{article.slug}.docx"
    doc.save(file_path)
    return file_path



def process_article(article: Article, client: genai.Client, settings: Settings) -> dict:
    translated_title = translate_text(client, article.title, "title")
    time.sleep(settings.request_sleep_seconds)
    translated_meta = translate_text(client, article.meta_description, "meta")
    time.sleep(settings.request_sleep_seconds)

    sections = split_html_into_sections(article.content_html, settings.section_char_limit)
    translated_sections: List[str] = []
    for index, section in enumerate(sections, start=1):
        print(f"    Translating section {index}/{len(sections)}")
        translated_sections.append(translate_text(client, section, "body_html"))
        time.sleep(settings.request_sleep_seconds)

    translated_html = "\n".join(translated_sections)
    shortcode_content = convert_html_to_shortcodes(translated_html)
    file_path = create_docx(article, translated_title, translated_meta, shortcode_content, settings.output_dir)

    return {
        "url": article.url,
        "slug": article.slug,
        "output_path": str(file_path),
        "sections": len(sections),
    }


def main() -> None:
    settings = load_settings()
    client = genai.Client(api_key=settings.gemini_api_key)
    gc = build_gspread_client(settings)
    sheet = gc.open_by_url(settings.sheet_url).worksheet(settings.sheet_name)
    pending_urls = get_pending_urls(sheet, settings)

    print(f"Found {len(pending_urls)} pending URLs")
    results = {"success": [], "failed": []}

    for index, url in enumerate(pending_urls, start=1):
        print(f"[{index}/{len(pending_urls)}] Processing: {url}")
        try:
            article = fetch_article(url, settings.request_timeout)
            if not article:
                raise ValueError("Could not find div.entry-content")
            result = process_article(article, client, settings)
            results["success"].append(result)
            print(f"  Saved: {result['output_path']}")
        except Exception as exc:
            print(f"  Failed: {exc}")
            results["failed"].append({"url": url, "error": str(exc)})

    summary_path = settings.output_dir / "run_summary.json"
    settings.output_dir.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Summary written to: {summary_path}")


if __name__ == "__main__":
    main()
